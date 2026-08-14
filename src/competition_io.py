from __future__ import annotations

from abc import ABC, abstractmethod
import csv
from dataclasses import asdict, dataclass, field
import json
from hashlib import sha256
from pathlib import Path
import re
from typing import Any, Iterable

from .adjudication import AdjudicationRecord, AthleteContext


@dataclass(slots=True)
class ImportIssue:
    row: int
    message: str


@dataclass(slots=True)
class ImportDraft:
    athletes: list[AthleteContext] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    source_description: str = ""

    @property
    def valid(self) -> bool:
        return bool(self.athletes) and not self.issues


@dataclass(slots=True)
class ExportReport:
    path: Path
    records_written: int
    warnings: list[str] = field(default_factory=list)


class CompetitionImportAdapter(ABC):
    @abstractmethod
    def parse(self, source: Path) -> ImportDraft:
        """Parse into an isolated draft. Never mutate an active session."""


class CompetitionExportAdapter(ABC):
    @abstractmethod
    def export(self, records: Iterable[AdjudicationRecord], destination: Path) -> ExportReport:
        """Export generic interchange without changing source records."""


_HEADER_ALIASES = {
    "name": {"name", "athlete", "competitor", "jmeno", "jméno", "zavodnik", "závodník"},
    "bib": {"bib", "bibnumber", "startnumber", "startovni cislo", "startovní číslo", "sc", "sč"},
    "club": {"club", "team", "oddil", "oddíl", "organisation", "organization"},
    "category": {"category", "group", "discipline", "kategorie", "skupina"},
    "start_order": {"startorder", "order", "poradi", "pořadí", "lane"},
    "external_id": {"id", "athleteid", "competitorid", "registrationid", "registrace"},
}


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9áčďéěíňóřšťúůýž]+", "", str(value).strip().lower())


def _canonical_headers(headers: Iterable[Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for raw in headers:
        normalised = _normalise_header(raw)
        for canonical, aliases in _HEADER_ALIASES.items():
            if normalised in {_normalise_header(alias) for alias in aliases}:
                result[canonical] = str(raw)
                break
    return result


def _draft_from_rows(rows: list[dict[str, Any]], description: str) -> ImportDraft:
    draft = ImportDraft(source_description=description)
    if not rows:
        draft.issues.append(ImportIssue(0, "The import contains no athlete rows"))
        return draft
    mapping = _canonical_headers(rows[0].keys())
    if "name" not in mapping:
        draft.issues.append(ImportIssue(1, "Map or provide an athlete name column"))
        return draft
    seen: set[str] = set()
    for index, row in enumerate(rows, start=2):
        name = str(row.get(mapping["name"], "") or "").strip()
        if not name:
            draft.issues.append(ImportIssue(index, "Athlete name is empty"))
            continue
        category = str(row.get(mapping.get("category", ""), "") or "Open").strip() or "Open"
        bib = str(row.get(mapping.get("bib", ""), "") or "").strip()
        club = str(row.get(mapping.get("club", ""), "") or "").strip()
        external_id = str(row.get(mapping.get("external_id", ""), "") or "").strip()
        try:
            start_order = int(float(str(row.get(mapping.get("start_order", ""), len(draft.athletes) + 1) or len(draft.athletes) + 1)))
        except ValueError:
            draft.issues.append(ImportIssue(index, "Start order is not a number"))
            continue
        identity = external_id or f"{category.casefold()}|{bib.casefold()}|{name.casefold()}|{club.casefold()}"
        if identity in seen:
            draft.issues.append(ImportIssue(index, "Duplicate athlete identity"))
            continue
        seen.add(identity)
        athlete_id = external_id or f"imported:{sha256(identity.encode('utf-8')).hexdigest()[:16]}"
        draft.athletes.append(AthleteContext(
            athlete_id=athlete_id,
            group=category,
            competitor_number=start_order,
            bib=bib,
            name=name,
            club=club,
            category=category,
            start_order=start_order,
            external_id=external_id,
        ))
    return draft


class CsvImportAdapter(CompetitionImportAdapter):
    def parse(self, source: Path) -> ImportDraft:
        raw = Path(source).read_bytes()
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1250"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            return ImportDraft(issues=[ImportIssue(0, "Unsupported text encoding")], source_description=str(source))
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.DictReader(text.splitlines(), dialect=dialect))
        return _draft_from_rows(rows, str(source))


class JsonImportAdapter(CompetitionImportAdapter):
    def parse(self, source: Path) -> ImportDraft:
        data = json.loads(Path(source).read_text(encoding="utf-8-sig"))
        if isinstance(data, dict):
            for key in ("athletes", "competitors", "entries", "roster"):
                if isinstance(data.get(key), list):
                    data = data[key]
                    break
        if not isinstance(data, list) or any(not isinstance(item, dict) for item in data):
            return ImportDraft(issues=[ImportIssue(0, "JSON must contain an array of athlete objects")], source_description=str(source))
        return _draft_from_rows(data, str(source))


class XlsxImportAdapter(CompetitionImportAdapter):
    def parse(self, source: Path) -> ImportDraft:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return ImportDraft(issues=[ImportIssue(0, "XLSX support requires openpyxl")], source_description=str(source))
        workbook = load_workbook(source, read_only=True, data_only=True)
        best_rows: list[dict[str, Any]] = []
        best_title = ""
        try:
            for sheet in workbook.worksheets:
                values = sheet.iter_rows(values_only=True)
                headers = next(values, None)
                if not headers:
                    continue
                rows = [dict(zip((str(value or "") for value in headers), row, strict=False)) for row in values if any(value not in (None, "") for value in row)]
                if "name" in _canonical_headers(headers) and len(rows) > len(best_rows):
                    best_rows, best_title = rows, sheet.title
        finally:
            workbook.close()
        return _draft_from_rows(best_rows, f"{source} [{best_title}]")


def adapter_for_path(path: Path) -> CompetitionImportAdapter:
    suffix = path.suffix.lower()
    if suffix == ".csv" or suffix in {".tsv", ".txt"}:
        return CsvImportAdapter()
    if suffix == ".json":
        return JsonImportAdapter()
    if suffix == ".xlsx":
        return XlsxImportAdapter()
    raise ValueError(f"Unsupported roster file type: {suffix or 'none'}")


def _export_row(record: AdjudicationRecord) -> dict[str, Any]:
    return {
        "record_id": record.record_id,
        "category": record.athlete.category or record.athlete.group,
        "bib": record.athlete.bib,
        "athlete": record.athlete.name,
        "club": record.athlete.club,
        "start_order": record.athlete.start_order,
        "attempt": record.attempt_number,
        "phase": record.competition_phase,
        "verdict": record.verdict,
        "distance_cm": str(record.distance_cm) if record.distance_cm is not None else "",
        "distance_source": record.distance_source,
        "wind_mps": f"{record.wind_tenths / 10:+.1f}" if record.wind_tenths is not None else "",
        "wind_source": record.wind_source,
        "verdict_timestamp": record.verdict_wall_time,
        "raw_evidence": record.evidence.raw_path,
        "annotated_evidence": record.evidence.annotated_path,
        "evidence_metadata": record.evidence.metadata_path,
        "clip": record.evidence.clip_path,
    }


class CsvExportAdapter(CompetitionExportAdapter):
    def export(self, records: Iterable[AdjudicationRecord], destination: Path) -> ExportReport:
        rows = [_export_row(record) for record in records]
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(destination.suffix + ".tmp")
        fields = list(rows[0]) if rows else list(_export_row(AdjudicationRecord("", 0, AthleteContext(""), 0, "", 0.0)))
        with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
        temporary.replace(destination)
        return ExportReport(destination, len(rows))


class JsonExportAdapter(CompetitionExportAdapter):
    def export(self, records: Iterable[AdjudicationRecord], destination: Path) -> ExportReport:
        rows = [_export_row(record) for record in records]
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(destination.suffix + ".tmp")
        temporary.write_text(json.dumps({"schema_version": 1, "records": rows}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        temporary.replace(destination)
        return ExportReport(destination, len(rows))


class ExperimentalAk2ExportAdapter(CompetitionExportAdapter):
    def export(self, records: Iterable[AdjudicationRecord], destination: Path) -> ExportReport:
        raise NotImplementedError(
            "AK2 export is intentionally disabled until an official sample is verified by a successful AK2 import"
        )
